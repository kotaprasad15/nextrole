"""Hand-built hostile files for the upload tests (no PDF library needed to *write* them, so the
test exercises our parser against bytes it did not produce itself)."""
import io
import zipfile

from docx import Document
from docx.shared import Pt, RGBColor

PAD = "Experienced backend engineer with Python, SQL and FastAPI. Built APIs and data pipelines. "


def text_op(text: str, x: float = 72, y: float = 700, size: float = 12, color: str = "", mode: int | None = None) -> str:
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    parts = ["q"]
    if color:
        parts.append(color)
    if mode is not None:
        parts.append(f"{mode} Tr")
    parts.append(f"BT /F1 {size} Tf {x} {y} Td ({safe}) Tj ET")
    parts.append("Q")
    return " ".join(parts)


def make_pdf(
    ops: str, *, catalog_extra: str = "", page_extra: str = "", extra_objects: tuple[str, ...] = (),
    info: str = "", mediabox: tuple[int, int, int, int] = (0, 0, 612, 792),
) -> bytes:
    """A one-page PDF. `catalog_extra`/`page_extra` inject raw dictionary entries; `extra_objects`
    are numbered from 7 (object 6 is the optional /Info dictionary)."""
    stream = ops.encode("latin-1")
    objs = [
        f"<< /Type /Catalog /Pages 2 0 R {catalog_extra} >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        f"<< /Type /Page /Parent 2 0 R /MediaBox [{' '.join(map(str, mediabox))}] "
        f"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> {page_extra} >>",
        None,  # content stream, written below
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< {info} >>",
        *extra_objects,
    ]
    out = bytearray(b"%PDF-1.7\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        if body is None:
            out += f"{i} 0 obj\n<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream\nendobj\n"
        else:
            out += f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    trailer = f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R"
    if info:
        trailer += " /Info 6 0 R"
    out += (trailer + f" >>\nstartxref\n{xref}\n%%EOF\n").encode()
    return bytes(out)


def plain_pdf(extra_ops: str = "") -> bytes:
    return make_pdf(text_op(PAD * 2) + " " + extra_ops)


def docx_bytes(visible: str = PAD * 2, *, hidden: str = "", white: str = "", tiny: str = "", author: str = "") -> bytes:
    doc = Document()
    p = doc.add_paragraph(visible)
    if hidden:
        doc.add_paragraph().add_run(hidden).font.hidden = True
    if white:
        run = doc.add_paragraph().add_run(white)
        run.font.color.rgb = RGBColor(255, 255, 255)
    if tiny:
        doc.add_paragraph().add_run(tiny).font.size = Pt(1)
    if author:
        doc.core_properties.author = author
        doc.core_properties.comments = "private note"
    del p
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def add_zip_member(data: bytes, name: str, payload: bytes) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item))
        dst.writestr(name, payload)
    return out.getvalue()
