"""Resume text extraction from PDF/DOCX with strict type and size checks.

Design decision: validate by file signature (magic bytes), not just the extension or the
client-supplied content type, since both are trivially spoofed.
"""
import io
import zipfile

from docx import Document
from pypdf import PdfReader


class ResumeError(ValueError):
    """User-correctable problem with an uploaded resume."""


def extract_text(data: bytes, max_bytes: int) -> str:
    if not data:
        raise ResumeError("empty file")
    if len(data) > max_bytes:
        raise ResumeError(f"file too large (max {max_bytes // (1024 * 1024)} MB)")
    if data.startswith(b"%PDF"):
        text = _pdf_text(data)
    elif data.startswith(b"PK") and zipfile.is_zipfile(io.BytesIO(data)):
        text = _docx_text(data)
    else:
        raise ResumeError("unsupported file type: upload a PDF or DOCX")
    if len(text.strip()) < 50:
        raise ResumeError("could not read text from the file (scanned image PDFs are unsupported)")
    return text


def _pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise ResumeError("could not parse PDF") from exc


def _docx_text(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeError("could not parse DOCX") from exc
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join(parts)
